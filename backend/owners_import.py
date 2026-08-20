"""
One-time / CLI import of Owners list Excel files into broker_contacts.

Email-extracted fields are kept. Excel fills blanks and extra columns.
Cells filled from Excel are tagged in excel_sourced (light-green in the UI).
"""
from __future__ import annotations

import io
import logging
import re
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from dotenv import load_dotenv

load_dotenv(Path(__file__).resolve().parent / ".env")

from agents.contact_extract import (  # noqa: E402
    CONTACT_FIELD_KEYS,
    _clean_contact_val,
    _merge_contact_fields,
    _parse_excel_sourced,
    _source_label,
    clean_company_name,
    contact_match_key,
    format_phone,
)
from column_defs import title_case_text  # noqa: E402
from database import supabase  # noqa: E402

logger = logging.getLogger("owners_import")

DEFAULT_FILE_1 = Path(r"C:\Users\Utkarsh\Downloads\Owners list - 1.xlsx")
DEFAULT_FILE_2 = Path(r"C:\Users\Utkarsh\Downloads\Owners list - 2.xls")

_EMAIL_RE = re.compile(r"[A-Z0-9._%+\-]+@[A-Z0-9.\-]+\.[A-Z]{2,}", re.I)
_URL_RE = re.compile(r"(https?://[^\s]+|www\.[^\s]+)", re.I)
_CITY_PAREN_RE = re.compile(r"^\((.+)\)$")
_COUNTRY_FIX = {
    "uk": "UK",
    "u.k.": "UK",
    "uae": "UAE",
    "usa": "USA",
    "us": "USA",
    "prc": "China",
    "s. korea": "Korea",
    "south korea": "Korea",
    "korea": "Korea",
    "p. r. china": "China",
    "p.r. china": "China",
    "pr china": "China",
}


def _cell(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    s = str(v).strip()
    if s.lower() in ("none", "nan", "-", "—", "–"):
        return ""
    return s


def _title(val: str) -> str:
    s = _cell(val)
    if not s:
        return ""
    return title_case_text(s) or s


def _country(val: str) -> str:
    s = _cell(val)
    if not s:
        return ""
    key = re.sub(r"\s+", " ", s.lower().strip(" ."))
    if key in _COUNTRY_FIX:
        return _COUNTRY_FIX[key]
    return _title(s)


def _emails(*vals: str) -> str:
    found: list[str] = []
    seen: set[str] = set()
    for raw in vals:
        for m in _EMAIL_RE.finditer(raw or ""):
            e = m.group(0).lower().rstrip(".,);")
            if e not in seen:
                seen.add(e)
                found.append(e)
    return ", ".join(found)


def _websites(*vals: str) -> str:
    found: list[str] = []
    seen: set[str] = set()
    for raw in vals:
        for m in _URL_RE.finditer(raw or ""):
            u = m.group(1).rstrip(".,);")
            key = u.lower()
            if key not in seen:
                seen.add(key)
                found.append(u)
    return ", ".join(found)


def _is_phone_value(raw: str) -> bool:
    s = _cell(raw)
    if not s or s.lower() in ("office", "fax", "tel", "telephone", "phone"):
        return False
    return len(re.sub(r"\D", "", s)) >= 5


_FLEET_TOKEN_RE = re.compile(
    r"""(?ix)^(?:
        \d+[kK]?|\d+\.\d+[kK]?|
        small|handy(?:size)?|mr|lr[12]?|panamax|aframax|afra|suezmax|vlcc|vlgc|
        cpp|dpp|chem(?:ical)?s?|palm(?:oil)?|tanker(?:s)?|vessel(?:s)?|vsls?|
        owner(?:s)?|operator(?:s)?|ownrs|fleet|dt|dwt\d*|kdwt|stst|ss|tanks|
        tc|t/?c|out|and|or|the|a|an|of|for|to|upto|up|mainly|pic|
        imo-?\d(?:/\d)?|west|med|east|one|more|built|cbm|mt|
        oil|chems|lpg|vess+els?|in
    )$"""
)
_NOTE_HINT_RE = re.compile(
    r"""(?ix)\b(?:
        j/v|joint\s+venture|bunker(?:ing)?|supplier|trader|player|
        managed\s+by|operated\s+by|time\s+charter|t/?c\s+out|tc\s+out|
        based\s+in|broking\s+company|distributor|subsidiary|spinoff|
        related\s+to|known\s+as|usually\s+|all\s+vessels|has\s+\d+\s+vessels|
        italian\s+owner|danish\s+owner|large\s+tanker\s+owner|
        palm\s+(?:oil\s+)?owner|palm\s+oil\s+to|dpp\s+owner|small\s+owner|small\s+shipowner|
        chartering\s+and\s+head|head\s+of\b|shipping\s+arm|
        only\s+do\s+their|sublet\s+to|manage\s+vessels|out\s+the\s+vessels|
        mainly\s+|mostly\s+in|for\s+domestic|for\s+international|
        unclear|pool\b|group\b|pte\.?|ltd\.?|gmbh|co\.,?|company\b|
        series\b|trading\s+in
    )\b"""
)
_COMPANYISH_RE = re.compile(
    r"(?i)\b(?:group|pte\.?|ltd\.?|gmbh|inc\.|llc|shipping arm|company ltd)\b"
)
_SERIES_RE = re.compile(r"(?i)\bseries\b")
_SPLIT_NAMES_RE = re.compile(r"\s*(?:,|;|&|\+| and )\s*", re.I)


def _norm_space(s: str) -> str:
    return re.sub(r"\s+", " ", s or "").strip(" \t,;/-")


def _is_fleet_category(text: str) -> bool:
    s = _norm_space(text)
    if not s:
        return True
    parts = [p for p in re.split(r"[\s,;/+&()\-]+", s) if p]
    if not parts:
        return True
    return all(_FLEET_TOKEN_RE.match(p) or p.replace(".", "", 1).isdigit() for p in parts)


def _clean_vessel_token(part: str) -> str:
    s = _norm_space(part)
    s = s.strip(").")
    s = re.sub(r"^\(?ex\s+", "", s, flags=re.I)
    s = re.sub(r"\s*\([^)]*\)\s*", " ", s)
    s = re.sub(r"\s*\([^)]*$", "", s)
    s = re.sub(r"\s*[-–]\s*\d[\d.]*\s*[kK]?\s*dwt\b.*$", "", s, flags=re.I)
    s = re.sub(r"\s*[-–]\s*IMO-?\d.*$", "", s, flags=re.I)
    s = re.sub(r"(?i)^(?:vesel|vessel)\s+name(?:\s+is)?\s*[:\-]?\s*", "", s)
    s = re.sub(r"(?i)^(?:vesel|vessel)\s+", "", s)
    s = _norm_space(s)
    if not s or s.lower() in ("unclear",) or _is_fleet_category(s) or _SERIES_RE.search(s):
        return ""
    if _COMPANYISH_RE.search(s) and not re.search(r"\d", s):
        return ""
    return s


def _join_vessels(names: list[str]) -> str:
    seen: set[str] = set()
    out: list[str] = []
    for n in names:
        key = n.lower()
        if n and key not in seen:
            seen.add(key)
            out.append(n)
    return "; ".join(out)


def _name_list(text: str) -> list[str]:
    parts = [p.strip() for p in _SPLIT_NAMES_RE.split(text) if p.strip()]
    names = [_clean_vessel_token(p) for p in parts]
    return [n for n in names if n]


def _other_if_extra(full: str, vessels: list[str]) -> str:
    """Keep remarks as other_info only when they add particulars beyond the names."""
    compact = full
    for v in vessels:
        compact = re.sub(re.escape(v), " ", compact, flags=re.I)
    compact = re.sub(
        r"(?i)\b(?:vesel|vessel)\s+name(?:\s+is)?|\b(?:head)?owners?\s+for\s+vessel|"
        r"owned chems tankers|operating \w+ vessel as|\bMT\b|\bex\b",
        " ",
        compact,
    )
    compact = _norm_space(re.sub(r"[\s,;:/()+\-]+", " ", compact))
    if not compact or _is_fleet_category(compact) or len(compact) <= 3:
        return ""
    return full


def split_remarks(raw: str) -> tuple[str, str]:
    """Split Excel remarks into (vessel_name, other_info). URLs are stripped."""
    text = _norm_space(raw)
    if not text:
        return "", ""
    leftover = _norm_space(_URL_RE.sub(" ", text)).strip(" -&,")
    leftover = _norm_space(re.sub(r"^(?:and|&)\s+|\s+(?:and|&)$", "", leftover, flags=re.I))
    if not leftover or leftover.lower() in ("and", "&"):
        return "", ""

    if _is_fleet_category(leftover) or _SERIES_RE.search(leftover):
        return "", leftover

    vessels: list[str] = []

    m = re.search(r"(?i)(?:vesel|vessel)\s+name(?:\s+is)?\s*[:\-]?\s*(.+)$", leftover)
    if m:
        tok = _clean_vessel_token(m.group(1))
        if tok:
            vessels.append(tok)
    else:
        m = re.fullmatch(r"(?i)vessel\s+([A-Z][A-Za-z0-9 .'-]+)", leftover)
        if m:
            tok = _clean_vessel_token(m.group(1))
            if tok:
                vessels.append(tok)

    m = re.search(r"(?i)(?:head)?owners?\s+for\s+vessel\s+([^(,]+)", leftover)
    if m:
        tok = _clean_vessel_token(m.group(1))
        if tok:
            vessels.append(tok)

    m = re.search(r"(?i)owned chems tankers\s+([A-Za-z][A-Za-z0-9 .'-]*)", leftover)
    if m:
        tok = _clean_vessel_token(m.group(1))
        if tok:
            vessels.append(tok)

    m = re.search(r"(?i)operating\s+\w+\s+vessel\s+as\s+(.+)$", leftover)
    if m:
        vessels.extend(_name_list(m.group(1)))

    m = re.search(
        r"(?i)vessel\s+\d[\w.]*\s*dwt\s*[-–]\s*([A-Za-z][A-Za-z0-9 .'-]+?)\s+trading",
        leftover,
    )
    if m:
        tok = _clean_vessel_token(m.group(1))
        if tok:
            vessels.append(tok)

    m = re.search(r"(?i)\bMT\s+([A-Za-z][A-Za-z0-9 .'-]*)", leftover)
    if m:
        tok = _clean_vessel_token(m.group(1))
        if tok:
            vessels.append(tok)

    for m in re.finditer(r"\b(Ever+ich\s+\d+)\b", leftover, re.I):
        vessels.append(_norm_space(m.group(1)))

    paren = re.search(r"\(([^)]*)", leftover)
    if paren:
        inner = paren.group(1)
        inner_names = _name_list(inner) if not _is_fleet_category(inner) else []
        if inner_names and (
            _is_fleet_category(leftover.split("(")[0]) or "owned chems" in leftover.lower()
        ):
            vessels.extend(inner_names)
        extra = re.findall(r"(?i)\b(Galissas)\b", leftover)
        vessels.extend(extra)

    if vessels:
        joined = _join_vessels(vessels)
        return joined, _other_if_extra(leftover, vessels)

    if _NOTE_HINT_RE.search(leftover) and not re.search(r"[,;]", leftover):
        return "", leftover
    if _COMPANYISH_RE.search(leftover) and not re.search(r"[,;]", leftover):
        return "", leftover

    paren = re.search(r"\(([^)]+)", leftover)
    if paren and _is_fleet_category(leftover.split("(")[0]):
        names = _name_list(paren.group(1))
        if names:
            return _join_vessels(names), leftover

    dash = re.match(r"^(.{2,40}?)\s*[-–]\s*(.+)$", leftover)
    if dash:
        left, right = _clean_vessel_token(dash.group(1)), dash.group(2).strip()
        if left and (
            right.lower() in ("spot",)
            or _is_fleet_category(right)
            or _NOTE_HINT_RE.search(right)
        ):
            extra = "" if _is_fleet_category(right) or right.lower() == "spot" else leftover
            return left, extra

    leftover_names = re.sub(r"\s*\(\s*ex\b[^)]*\)?", "", leftover, flags=re.I)
    leftover_names = _norm_space(leftover_names)

    names = _name_list(leftover_names)
    parts = [p.strip() for p in _SPLIT_NAMES_RE.split(leftover_names) if p.strip()]
    if names and len(names) == len(parts):
        return _join_vessels(names), ""
    if names and not _is_fleet_category(leftover):
        return _join_vessels(names), leftover if _NOTE_HINT_RE.search(leftover) else ""
    return "", leftover


def _remarks_fields(remarks: str) -> tuple[str, str, str]:
    """Returns vessel_name, other_info, websites-from-remarks."""
    vessel_name, other_info = split_remarks(remarks)
    return vessel_name, other_info, _websites(remarks)


def _empty_row() -> dict[str, str]:
    return {k: "" for k in CONTACT_FIELD_KEYS}


def _person_row(**kwargs: str) -> dict[str, str]:
    row = _empty_row()
    row["status"] = "Active"
    for k, v in kwargs.items():
        if k in row:
            row[k] = _cell(v)
    row["company"] = clean_company_name(row.get("company") or "")
    row["contact_name"] = _title(row.get("contact_name") or "")
    row["country"] = _country(row.get("country") or "")
    row["city"] = _title(row.get("city") or "")
    row["email"] = _emails(row.get("email") or "")
    row["website_address"] = _websites(row.get("website_address") or "")
    row["off_phone"] = format_phone(row.get("off_phone") or "")
    row["mob_phone"] = format_phone(row.get("mob_phone") or "")
    row["fax"] = format_phone(row.get("fax") or "")
    row["office_address"] = re.sub(r"\s+", " ", row.get("office_address") or "").strip()
    row["trade"] = re.sub(r"\s+", " ", row.get("trade") or "").strip()
    row["role"] = _normalize_role(row.get("role") or "")
    row["vessel_name"] = re.sub(r"\s+", " ", row.get("vessel_name") or "").strip()
    row["other_info"] = re.sub(r"\s+", " ", row.get("other_info") or "").strip()
    return row


def _normalize_role(raw: str) -> str:
    s = _cell(raw).strip(" /")
    if not s:
        return ""
    low = s.lower()
    parts = []
    if "owner" in low:
        parts.append("Owner")
    if "broker" in low:
        parts.append("Broker")
    if "operator" in low:
        parts.append("Operator")
    if "charter" in low:
        parts.append("Charterer")
    return "/".join(parts) if parts else _title(s)


def _has_contact_signal(row: dict[str, str]) -> bool:
    return bool(
        row.get("contact_name")
        or row.get("email")
        or row.get("mob_phone")
        or row.get("off_phone")
        or row.get("company")
    )


def parse_owners_xlsx_matrix(rows: list) -> list[dict[str, str]]:
    """Icon Owners list workbook: company group + PIC rows (xlsx layout)."""
    if not rows:
        return []
    company = address = country = office_tel = office_email = website = trade = remarks = ""
    out: list[dict[str, str]] = []
    start = 1 if rows and any(_cell(c) for c in (rows[0] or [])) else 0
    for raw in rows[start:]:
        r = list(raw or []) + [None] * 12
        co = _cell(r[1])
        if co:
            company = clean_company_name(co)
            address = country = office_tel = office_email = website = trade = remarks = ""
        if _cell(r[2]):
            address = _cell(r[2])
        if _cell(r[3]):
            country = _cell(r[3])
        if _is_phone_value(r[4]):
            office_tel = _cell(r[4])
        pic = _cell(r[5])
        mobile = _cell(r[6])
        pic_email_cell = _cell(r[7])
        office_email_cell = _cell(r[8])
        web_cell = _cell(r[9])
        if _cell(r[10]):
            trade = _cell(r[10])
        if _cell(r[11]):
            remarks = _cell(r[11])

        emails = _emails(pic_email_cell, office_email_cell, web_cell, office_email)
        if _emails(office_email_cell):
            office_email = office_email_cell
        vessel_name, other_info, remark_sites = _remarks_fields(remarks)
        websites = _websites(web_cell, pic_email_cell, office_email_cell, website, remark_sites)
        if _websites(web_cell):
            website = web_cell

        row = _person_row(
            company=company,
            contact_name=pic,
            email=emails,
            off_phone=office_tel,
            mob_phone=mobile,
            website_address=websites,
            office_address=address,
            country=country,
            trade=trade,
            vessel_name=vessel_name,
            other_info=other_info,
            role="Owner",
        )
        if pic or emails or mobile:
            if _has_contact_signal(row):
                out.append(row)
    return out


def parse_owners_list_1(path: Path) -> list[dict[str, str]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    return parse_owners_xlsx_matrix(rows)


def parse_owners_xlsx_bytes(raw: bytes) -> list[dict[str, str]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    out: list[dict[str, str]] = []
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        out.extend(parse_owners_xlsx_matrix(rows))
    wb.close()
    return out


def parse_owners_xls_book(book) -> list[dict[str, str]]:
    """Older Owners list .xls: (City) group lines + customer/PIC rows."""
    out: list[dict[str, str]] = []
    for sheet_name in book.sheet_names():
        sh = book.sheet_by_name(sheet_name)
        if sh.nrows < 4:
            continue
        company = city = role = office = fax = remarks = ""
        for r in range(3, sh.nrows):
            vals = [_cell(sh.cell_value(r, c)) for c in range(sh.ncols)]
            if not any(vals):
                continue
            pad = vals + [""] * 10
            cust, status, office_c, fax_c, pic, mobile, messenger, email_c, rem = (
                pad[1], pad[2], pad[3], pad[4], pad[5], pad[6], pad[7], pad[8], pad[9]
            )
            city_m = _CITY_PAREN_RE.match(cust)
            if city_m:
                city = city_m.group(1).strip()
                continue
            if cust:
                company = clean_company_name(cust)
                city = role = office = fax = remarks = ""
            if status and status.lower() not in ("status",):
                role = status
            if _is_phone_value(office_c):
                office = office_c
            if _is_phone_value(fax_c) or (fax_c and re.search(r"\d", fax_c) and fax_c.lower() not in ("fax",)):
                fax = fax_c
            if rem:
                remarks = rem
            emails = _emails(email_c)
            vessel_name, other_info, remark_sites = _remarks_fields(remarks)
            websites = _websites(email_c, remark_sites)
            if messenger and messenger.lower() not in ("messenger",):
                other_info = f"{other_info}; {messenger}".strip("; ")
            row = _person_row(
                company=company,
                contact_name=pic,
                email=emails,
                off_phone=office,
                mob_phone=mobile,
                fax=fax,
                website_address=websites,
                city=city,
                role=role,
                vessel_name=vessel_name,
                other_info=other_info,
            )
            if pic or emails or mobile:
                if _has_contact_signal(row):
                    out.append(row)
    return out


def parse_owners_xls(path: Path) -> list[dict[str, str]]:
    import xlrd

    book = xlrd.open_workbook(str(path))
    return parse_owners_xls_book(book)


def parse_owners_xls_bytes(raw: bytes) -> list[dict[str, str]]:
    import xlrd

    book = xlrd.open_workbook(file_contents=raw)
    return parse_owners_xls_book(book)


def _email_set(raw: str) -> set[str]:
    return {
        e.strip().lower()
        for e in re.split(r"\s*[,;]\s*", raw or "")
        if e.strip() and "@" in e
    }


def _index_existing(rows: list[dict[str, Any]]) -> tuple[dict[str, dict], dict[str, dict], dict[str, dict]]:
    by_key: dict[str, dict] = {}
    by_email: dict[str, dict] = {}
    by_name_co: dict[str, dict] = {}
    for row in rows:
        key = row.get("match_key") or contact_match_key(row)
        if key and key not in by_key:
            by_key[key] = row
        for e in _email_set(row.get("email") or ""):
            by_email.setdefault(e, row)
        name = (row.get("contact_name") or "").strip().lower()
        co = (row.get("company") or "").strip().lower()
        if name and co:
            by_name_co.setdefault(f"{name}|{co}", row)
    return by_key, by_email, by_name_co


def _find_match(
    contact: dict[str, str],
    by_key: dict[str, dict],
    by_email: dict[str, dict],
    by_name_co: dict[str, dict],
) -> dict | None:
    key = contact_match_key(contact)
    if key and key in by_key:
        return by_key[key]
    name = (contact.get("contact_name") or "").strip()
    if name:
        for e in _email_set(contact.get("email") or ""):
            hit = by_email.get(e)
            if hit and (hit.get("contact_name") or "").strip().lower() == name.lower():
                return hit
        nk = f"{name.lower()}|{(contact.get('company') or '').strip().lower()}"
        if nk in by_name_co:
            return by_name_co[nk]
    else:
        for e in _email_set(contact.get("email") or ""):
            if e in by_email:
                return by_email[e]
    return None


OWNERS_REPLACE_KEYS = {
    "vessel_name",
    "other_info",
    "off_phone",
    "fax",
    "website_address",
    "trade",
    "office_address",
    "country",
    "city",
    "role",
}


def _is_owners_only(row: dict[str, Any]) -> bool:
    src = str(row.get("source") or "").lower()
    if src == "owners-list":
        return True
    excel = _parse_excel_sourced(row.get("excel_sourced"))
    return bool(excel) and src not in ("email", "mixed")


def _fill_from_excel(
    existing: dict[str, Any],
    incoming: dict[str, str],
    *,
    overwrite_excel_fields: bool = False,
) -> tuple[dict[str, Any], list[str]]:
    """Keep existing values; fill blanks from Excel; union emails/phones.

    When overwrite_excel_fields is True (owners-list-only rows), replace
    leaky Excel columns even if the current DB value is non-empty.
    """
    sourced = set(_parse_excel_sourced(existing.get("excel_sourced")))
    merged = _merge_contact_fields(existing, {})  # copy cleaned
    for key in CONTACT_FIELD_KEYS:
        cur = _clean_contact_val(existing.get(key))
        nxt = _clean_contact_val(incoming.get(key))
        if key == "status":
            continue
        if overwrite_excel_fields and key in OWNERS_REPLACE_KEYS:
            merged[key] = nxt
            if nxt:
                sourced.add(key)
            else:
                sourced.discard(key)
            continue
        if key in ("email", "company"):
            if nxt and not cur:
                merged[key] = nxt
                sourced.add(key)
            elif nxt and cur:
                from agents.contact_extract import _merge_text_field

                merged[key] = _merge_text_field(cur, nxt)
            else:
                merged[key] = cur
            continue
        if key in ("off_phone", "mob_phone", "fax"):
            if nxt and not cur:
                merged[key] = nxt
                sourced.add(key)
            elif nxt and cur:
                from agents.contact_extract import _merge_text_field

                merged[key] = _merge_text_field(cur, nxt, sep=" ; ")
            else:
                merged[key] = cur
            continue
        if nxt and not cur:
            merged[key] = nxt
            sourced.add(key)
        else:
            merged[key] = cur
    vn = _clean_contact_val(merged.get("vessel_name"))
    oi = _clean_contact_val(merged.get("other_info"))
    if vn and oi and vn == oi:
        merged["other_info"] = ""
        sourced.discard("other_info")
    return merged, sorted(sourced)


def import_owner_contacts(parsed: list[dict[str, str]]) -> dict[str, int]:
    """Upsert owners-directory rows. Never writes into email-extracted contacts."""
    existing = (
        supabase.table("broker_contacts")
        .select("id, match_key, attachment_id, parent_email_id, excel_sourced, source, "
                + ", ".join(CONTACT_FIELD_KEYS))
        .execute()
    ).data or []
    existing = [r for r in existing if _is_owners_only(r)]
    by_key, by_email, by_name_co = _index_existing(existing)

    inserted = 0
    updated = 0
    skipped = 0
    now = datetime.now(timezone.utc).isoformat()

    for incoming in parsed:
        if not _has_contact_signal(incoming):
            skipped += 1
            continue
        hit = _find_match(incoming, by_key, by_email, by_name_co)
        if hit:
            overwrite = _is_owners_only(hit)
            merged, sourced = _fill_from_excel(
                hit, incoming, overwrite_excel_fields=overwrite
            )
            has_email_parent = bool(hit.get("parent_email_id"))
            patch = dict(merged)
            patch["excel_sourced"] = sourced
            patch["source"] = _source_label(sourced, has_email_parent=has_email_parent)
            patch["match_key"] = contact_match_key(merged) or hit.get("match_key") or ""
            patch["updated_at"] = now
            supabase.table("broker_contacts").update(patch).eq("id", hit["id"]).execute()
            hit.update(patch)
            updated += 1
            continue

        sourced = [k for k in CONTACT_FIELD_KEYS if _clean_contact_val(incoming.get(k)) and k != "status"]
        payload = dict(incoming)
        payload["match_key"] = contact_match_key(incoming)
        payload["excel_sourced"] = sourced
        payload["source"] = "owners-list"
        payload["used_fallback"] = False
        payload["updated_at"] = now
        result = supabase.table("broker_contacts").insert(payload).execute()
        row = result.data[0] if result.data else payload
        key = payload["match_key"]
        if key:
            by_key[key] = row
        for e in _email_set(incoming.get("email") or ""):
            by_email.setdefault(e, row)
        name = (incoming.get("contact_name") or "").strip().lower()
        co = (incoming.get("company") or "").strip().lower()
        if name and co:
            by_name_co.setdefault(f"{name}|{co}", row)
        inserted += 1

    return {
        "parsed_total": len(parsed),
        "inserted": inserted,
        "updated": updated,
        "skipped": skipped,
    }


def import_owners(file1: Path, file2: Path) -> dict[str, int]:
    # Older workbook first, then the newer PIC list so overlaps keep file-1 values.
    parsed: list[dict[str, str]] = []
    n2 = 0
    if file2.exists():
        extra = parse_owners_xls(file2)
        parsed.extend(extra)
        n2 = len(extra)
        logger.info("Parsed %s → %d PIC rows", file2.name, n2)
    else:
        logger.warning("Missing %s", file2)
    n1 = 0
    if file1.exists():
        extra = parse_owners_list_1(file1)
        parsed.extend(extra)
        n1 = len(extra)
        logger.info("Parsed %s → %d PIC rows", file1.name, n1)
    else:
        logger.warning("Missing %s", file1)

    stats = import_owner_contacts(parsed)
    stats["parsed_file1"] = n1
    stats["parsed_file2"] = n2
    logger.info("Owners import %s", stats)
    return stats


def main() -> None:
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
    f1 = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_FILE_1
    f2 = Path(sys.argv[2]) if len(sys.argv) > 2 else DEFAULT_FILE_2
    print(import_owners(f1, f2))


if __name__ == "__main__":
    main()
